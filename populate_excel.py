import urllib.request
import urllib.parse
import xml.etree.ElementTree as ET
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def search_arxiv(query, max_results=5):
    url = f'http://export.arxiv.org/api/query?search_query={urllib.parse.quote(query)}&start=0&max_results={max_results}&sortBy=relevance&sortOrder=descending'
    response = urllib.request.urlopen(url).read()
    root = ET.fromstring(response)
    ns = {'atom': 'http://www.w3.org/2005/Atom'}
    
    papers = []
    for entry in root.findall('atom:entry', ns):
        title = entry.find('atom:title', ns).text.strip().replace('\n', ' ')
        authors = ", ".join([a.find('atom:name', ns).text for a in entry.findall('atom:author', ns)])
        published = entry.find('atom:published', ns).text[:4]
        
        # Look for doi
        doi = "Non rapporté dans la publication"
        for link in entry.findall('atom:link', ns):
            if link.attrib.get('title') == 'doi':
                doi = link.attrib.get('href')
                
        link = entry.find('atom:id', ns).text
        papers.append({
            "Référence": f"[{len(papers)+1}]",
            "Titre": title,
            "Auteurs": authors,
            "Année": published,
            "Revue/Conférence": "arXiv preprint",
            "DOI": doi,
            "Type de publication": "Prépublication (arXiv)",
            "Évaluation par les pairs": "Non vérifié",
            "Lien officiel": link,
            "Sous-thème associé": "Modèles / Edge",
            "Métadonnées": f"{authors}, '{title}', {published}."
        })
    return papers

print("Recherche de publications sur arXiv...")
papers_edge = search_arxiv('all:"anomaly detection" AND all:"edge computing"', 5)
papers_xai = search_arxiv('all:"anomaly detection" AND all:"explainable" OR all:"SHAP"', 5)

all_papers = papers_edge + papers_xai
for i, p in enumerate(all_papers):
    p["Référence"] = f"[{i+1}]"

# Load the excel and fill the "1_Sources" sheet
file_path = r'c:\Users\marie\Desktop\Projets\Stage antigravity\Etat_de_l_art_Bibliographie.xlsx'
df_sources = pd.DataFrame(all_papers)

print("Mise à jour du fichier Excel...")
with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
    columns = ["Référence", "Métadonnées", "DOI", "Type de publication", "Évaluation par les pairs", "Lien officiel", "Sous-thème associé"]
    df_sources_formatted = df_sources[columns]
    df_sources_formatted.to_excel(writer, sheet_name="1_Sources", index=False)

# Format the Excel using openpyxl
wb = openpyxl.load_workbook(file_path)
header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    # Apply header style
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Adjust column widths and text wrap
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
            # Set alignment with wrap text
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            
        adjusted_width = (max_length + 2)
        if adjusted_width > 50:
            adjusted_width = 50
        ws.column_dimensions[column].width = adjusted_width

wb.save(file_path)
print(f"Excel formaté professionnellement et pré-rempli avec {len(all_papers)} publications sources.")
