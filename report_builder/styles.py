"""
report_builder/styles.py
------------------------
Gestion des styles typographiques, des couleurs de charte graphique,
des marges, des tableaux stylisés et des encadrés d'expertise.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Charte Graphique "Executive / Academic"
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Titres principaux, en-têtes
COLOR_STEEL = RGBColor(43, 84, 126)     # #2B547E - Titres secondaires
COLOR_CHARCOAL = RGBColor(45, 55, 72)   # #2D3748 - Sous-titres
COLOR_BODY = RGBColor(34, 34, 34)       # #222222 - Corps de texte
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B - Légendes, métadonnées
COLOR_GREEN = RGBColor(46, 125, 50)     # #2E7D32 - Succès / Champion Edge
COLOR_RED = RGBColor(198, 40, 40)       # #C62828 - Alertes / Anomalies

HEX_NAVY = "1B365D"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_ALT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_BORDER_DARK = "1B365D"

def setup_document_styles(doc):
    """Configure les marges, polices de base et en-têtes/pieds de page."""
    # Marges A4 2.5 cm (standard universitaire)
    for section in doc.sections:
        section.top_margin = Inches(0.98)     # ~2.5 cm
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)
        section.page_width = Inches(8.27)     # A4
        section.page_height = Inches(11.69)
        
        # En-tête
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run("Rapport de Stage de Fin d'Études — Edge AI & RAG sur Raspberry Pi")
        h_run.font.name = 'Calibri'
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = COLOR_MUTED
        
        # Pied de page
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = footer_p.add_run("ESPRIM — Master / Diplôme National d'Ingénieur  |  Page ")
        f_run.font.name = 'Calibri'
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = COLOR_MUTED
        
        # Ajouter le champ de numéro de page XML
        f_field = footer_p.add_run()
        f_field.font.name = 'Calibri'
        f_field.font.size = Pt(8.5)
        f_field.font.bold = True
        f_field.font.color.rgb = COLOR_NAVY
        f_xml = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        f_field._r.append(f_xml)

    # Configuration des styles par défaut
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = COLOR_BODY
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def add_heading_1(doc, text, page_break=True):
    """Titre de niveau 1 (Chapitre) avec saut de page et style élégant."""
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_2(doc, text):
    """Titre de niveau 2 (Sous-chapitre)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_STEEL
    return p

def add_heading_3(doc, text):
    """Titre de niveau 3."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_CHARCOAL
    return p

def add_paragraph(doc, text, bold_prefix=None, space_after=5):
    """Ajoute un paragraphe justifié avec un éventuel préfixe en gras."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Calibri'
        r_bold.font.size = Pt(11)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_NAVY
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = COLOR_BODY
    return p

def add_bullet_point(doc, text, bold_title=None):
    """Ajoute un point de liste avec puce personnalisée."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    if bold_title:
        r_bold = p.add_run(bold_title + " : ")
        r_bold.font.name = 'Calibri'
        r_bold.font.size = Pt(11)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_STEEL
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = COLOR_BODY
    return p

def add_callout_box(doc, title, text, box_type="info"):
    """Crée un encadré stylisé avec fond teinté et bordure gauche épaisse."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    
    # Couleur de bordure et de titre selon le type
    if box_type == "decision":
        border_hex = "2E7D32" # Vert
        title_color = COLOR_GREEN
        prefix_icon = "[DÉCISION CLÉ DE CONCEPTION] "
    elif box_type == "alert":
        border_hex = "C62828" # Rouge
        title_color = COLOR_RED
        prefix_icon = "[ALERTE SÉCURITÉ / INDUSTRIE] "
    elif box_type == "norme":
        border_hex = "1B365D" # Bleu nuit
        title_color = COLOR_NAVY
        prefix_icon = "[NORME & STANDARD INDUSTRIEL] "
    else:
        border_hex = "2B547E" # Bleu acier
        title_color = COLOR_STEEL
        prefix_icon = "[REMARQUE D'EXPERTISE] "

    # XML Shading (Fond) et Borders (Bordure gauche 32 = 4pt)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), HEX_LIGHT_BLUE))
    borders = parse_xml(r'''
        <w:tcBorders %s>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="32" w:space="0" w:color="%s"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''' % (nsdecls('w'), border_hex))
    tcPr.append(shd)
    tcPr.append(borders)
    
    # Marges intérieures de la cellule (padding)
    tcMar = parse_xml(r'''
        <w:tcMar %s>
            <w:top w:w="140" w:type="dxa"/>
            <w:left w:w="200" w:type="dxa"/>
            <w:bottom w:w="140" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    ''' % nsdecls('w'))
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    run_title = p.add_run(prefix_icon + title + "\n")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = title_color
    
    run_text = p.add_run(text)
    run_text.font.name = 'Calibri'
    run_text.font.size = Pt(10)
    run_text.font.italic = False
    run_text.font.color.rgb = COLOR_BODY
    
    # Espace après le tableau
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(doc, image_path, figure_number, caption, width_inches=5.8):
    """Insère une figure centrée avec une légende numérotée académique."""
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    
    try:
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        
        r_num = p_cap.add_run(f"Figure {figure_number} — ")
        r_num.font.name = 'Calibri'
        r_num.font.size = Pt(9.5)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_NAVY
        
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED
    except Exception as e:
        print(f"Warning: Could not insert image {image_path}: {e}")

def create_styled_table(doc, headers, rows_data, col_widths=None, table_title=None, table_number=None):
    """Génère un tableau élégant avec en-tête bleu nuit, lignes alternées et bordures fines."""
    if table_title and table_number:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(3)
        p_title.paragraph_format.keep_with_next = True
        
        r_num = p_title.add_run(f"Tableau {table_number} — ")
        r_num.font.name = 'Calibri'
        r_num.font.size = Pt(10)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_NAVY
        
        r_tit = p_title.add_run(table_title)
        r_tit.font.name = 'Calibri'
        r_tit.font.size = Pt(10)
        r_tit.font.italic = True
        r_tit.font.color.rgb = COLOR_BODY

    n_rows = len(rows_data) + 1
    n_cols = len(headers)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 1. Formatage de l'en-tête (Ligne 0)
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    for c_idx, head_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), HEX_NAVY)))
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(head_text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # 2. Formatage des lignes de données
    for r_idx, row_values in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        fill_hex = HEX_ALT_ROW if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill_hex)))
            
            # Bordures fines horizontales
            tcBorders = parse_xml(r'''
                <w:tcBorders %s>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="%s"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="%s"/>
                    <w:left w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            ''' % (nsdecls('w'), HEX_BORDER, HEX_BORDER))
            tcPr.append(tcBorders)

            p = cell.paragraphs[0]
            # Aligner à gauche pour la 1ère colonne, au centre pour les métriques numériques
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            
            # Détecter si valeur championne ou clé
            r = p.add_run(str(val))
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if "★" in str(val) or "1.0000" in str(val) or "Gradient Boosting" in str(val):
                r.font.bold = True
                r.font.color.rgb = COLOR_NAVY
            else:
                r.font.color.rgb = COLOR_BODY

    # Appliquer largeurs de colonnes si spécifiées
    if col_widths and len(col_widths) == n_cols:
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table
