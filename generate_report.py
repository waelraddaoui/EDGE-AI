from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_report():
    doc = Document()
    
    # Titre
    title = doc.add_heading("Point d'Avancement : Projet Stage Ingénieur (Edge AI & Maintenance Prédictive)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("").add_run("Date : Août 2026").italic = True
    
    # Section 1 : Introduction et Objectifs
    doc.add_heading("1. Objectifs du Projet", level=1)
    p = doc.add_paragraph(
        "Ce projet s'inscrit dans le cadre d'un stage de fin d'études ingénieur visant à développer une "
        "solution complète d'Intelligence Artificielle de bout en bout pour la maintenance prédictive, "
        "spécifiquement orientée vers le déploiement sur systèmes embarqués (Edge AI / Raspberry Pi)."
    )
    
    # Section 2 : Travail Réalisé (Pipeline)
    doc.add_heading("2. Travail Réalisé (Pipeline Global)", level=1)
    doc.add_paragraph("L'architecture complète du projet a été implémentée avec succès. Les étapes suivantes sont désormais fonctionnelles :", style='BodyText')
    
    doc.add_paragraph("Acquisition et Traitement du Signal (Jeu de données CWRU) : Extraction de caractéristiques vibratoires temporelles et fréquentielles (FFT, Kurtosis, RMS).", style='List Bullet')
    doc.add_paragraph("Entraînement des modèles IA : Entraînement de plusieurs architectures pour la classification (sain vs défaut).", style='List Bullet')
    doc.add_paragraph("Développement d'une API Backend (FastAPI) : Pont entre les algorithmes d'IA et l'utilisateur.", style='List Bullet')
    doc.add_paragraph("Création d'un Dashboard Professionnel (React / Vite) : Interface utilisateur moderne intégrant une simulation matérielle temps réel et un système de chat.", style='List Bullet')
    doc.add_paragraph("Système RAG Intégré (Retrieval-Augmented Generation) : Assistant virtuel expert capable de répondre aux questions techniques en s'appuyant sur la recherche sémantique.", style='List Bullet')
    
    # Section 3 : Modèles Développés
    doc.add_heading("3. Modèles d'IA Développés et Justification", level=1)
    
    doc.add_paragraph("A. Modèle Machine Learning (Gradient Boosting)")
    p_rf = doc.add_paragraph()
    p_rf.add_run("Objectif : ").bold = True
    p_rf.add_run("Preuve de concept rapide et déploiement Edge.\n")
    p_rf.add_run("Avantages : ").bold = True
    p_rf.add_run("Exécution ultra-rapide sur processeur ARM (Raspberry Pi), faible consommation mémoire. Excellent pour les données statistiques extraites (Kurtosis, etc.).")
    
    doc.add_paragraph("B. Modèles Deep Learning (PyTorch : CNN 1D & Autoencodeur)")
    p_dl = doc.add_paragraph()
    p_dl.add_run("Objectif : ").bold = True
    p_dl.add_run("Capturer les dépendances temporelles complexes directement sur le signal brut.\n")
    p_dl.add_run("Statut : ").bold = True
    p_dl.add_run("Les modèles ont été entraînés et exportés. La prochaine étape physique consiste à les convertir en ONNX pour l'inférence matérielle.")
    
    doc.add_paragraph("C. Modèle de Traitement du Langage (Sentence-Transformers)")
    p_nlp = doc.add_paragraph()
    p_nlp.add_run("Objectif : ").bold = True
    p_nlp.add_run("Moteur de l'assistant RAG.\n")
    p_nlp.add_run("Technologie : ").bold = True
    p_nlp.add_run("Modèle dense 'paraphrase-multilingual-MiniLM-L12-v2' pour permettre à l'opérateur de poser des questions à la documentation technique via recherche sémantique.")

    # Section 4 : Tests et Validation
    doc.add_heading("4. Tests et Validation des Performances", level=1)
    doc.add_paragraph("Un système de test global automatisé (run_all_tests.py) a été mis en place pour certifier le bon fonctionnement du pipeline. Tous les tests sont actuellement passés avec succès (Exit Code 0).")
    doc.add_paragraph("- Test 1 : Validation des données CWRU (OK)", style='List Bullet')
    doc.add_paragraph("- Test 2 : Extraction des caractéristiques vibratoires (OK)", style='List Bullet')
    doc.add_paragraph("- Test 3 : Inférence du modèle Gradient Boosting (OK)", style='List Bullet')
    doc.add_paragraph("- Test 4 : Backend API (FastAPI) fonctionnel (OK)", style='List Bullet')
    doc.add_paragraph("- Test 5 : Recherche sémantique du RAG fonctionnelle (OK)", style='List Bullet')
    
    # Section 5 : Prochaines Étapes
    doc.add_heading("5. Prochaines Étapes", level=1)
    doc.add_paragraph("1. Déploiement physique : Exporter les modèles PyTorch en ONNX et tester l'inférence réelle sur le processeur du Raspberry Pi.", style='List Number')
    doc.add_paragraph("2. Ajustement fin (Fine-tuning) : Utilisation potentielle d'Optuna pour maximiser la précision algorithmique des modèles.", style='List Number')
    doc.add_paragraph("3. Préparation de la soutenance : Formalisation des résultats dans le rapport de stage final.", style='List Number')
    
    # Sauvegarde
    output_path = os.path.join(os.getcwd(), "Rapport_Avancement_Stage.docx")
    doc.save(output_path)
    print(f"Rapport généré : {output_path}")

if __name__ == "__main__":
    create_report()
